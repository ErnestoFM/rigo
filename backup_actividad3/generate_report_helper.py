"""
================================================================================
ACTIVIDAD 3: ANÁLISIS SINTÁCTICO DE VARIABLES Y CONSTANTES
MATERIA: TRADUCTORES DE LENGUAJE / COMPILADORES
PROFESOR: RIGOBERTO CARDENAS LARIOS

INTEGRANTES DEL EQUIPO (ORDEN ALFABÉTICO POR PRIMER APELLIDO):
1. DIAZ TORRES JAZMIN MONTSERRAT
2. FIERRO MELÉNDEZ ERNESTO HATUEY
3. HERNANDEZ GARCIA HECTOR GABRIEL
4. MILLAN GUERRERO OSWALDO JOSUE
================================================================================
"""

import os
import sys
from io import StringIO
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import main

DOCX_NAME = "DiazTorresJazmin_FierroMelendezErnesto_HernandezGarciaHector_MillanGuerreroOswaldo_A3_AnalisisSintacticoVariablesConstantes.docx"
TXT_OUTPUT = "Salida_Ejecucion_Consola.txt"

def set_cell_background(cell, fill_hex="F1F5F9"):
    """Aplica color de fondo a una celda de tabla en python-docx."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Establece márgenes internos de celda en dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def export_execution_log() -> str:
    """Ejecuta main.py capturando stdout y guardando en Salida_Ejecucion_Consola.txt."""
    print("⏳ Ejecutando suite de pruebas sintácticas de main.py...")
    buffer = StringIO()
    old_stdout = sys.stdout
    sys.stdout = buffer
    
    try:
        main.main()
    finally:
        sys.stdout = old_stdout
        
    output = buffer.getvalue()
    
    with open(TXT_OUTPUT, "w", encoding="utf-8") as f:
        f.write(output)
        
    print(f"✔ Salida de ejecución guardada en '{TXT_OUTPUT}' ({len(output)} caracteres).")
    return output

def add_code_block(doc, code_text: str, bg_hex="F8FAFC"):
    """Agrega un bloque de código formateado con caja y tipografía Consolas."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph() # Espacio posterior

def generate_report_docx(console_output: str):
    """Crea el documento oficial de entrega con formato riguroso."""
    print(f"📄 Generando reporte oficial en '{DOCX_NAME}'...")
    doc = docx.Document()
    
    # Configuración de márgenes estándar (1 pulgada)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # ---------------- PORTADA ----------------
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(18)
    
    run_inst = title_p.add_run("MATERIA: TRADUCTORES DE LENGUAJE / COMPILADORES\nPROFESOR: RIGOBERTO CÁRDENAS LARIOS\n\n")
    run_inst.font.name = "Arial"
    run_inst.font.size = Pt(13)
    run_inst.font.bold = True
    run_inst.font.color.rgb = RGBColor(15, 23, 42)
    
    run_act = title_p.add_run("ACTIVIDAD 3: ANÁLISIS SINTÁCTICO DE VARIABLES Y CONSTANTES\nDESARROLLO DE UN COMPILADOR EN PYTHON\n")
    run_act.font.name = "Arial"
    run_act.font.size = Pt(15)
    run_act.font.bold = True
    run_act.font.color.rgb = RGBColor(30, 64, 175)

    integrantes_p = doc.add_paragraph()
    integrantes_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    integrantes_p.paragraph_format.space_after = Pt(14)
    
    run_int_title = integrantes_p.add_run("INTEGRANTES DEL EQUIPO\n(Orden Alfabético por Primer Apellido):\n\n")
    run_int_title.font.name = "Arial"
    run_int_title.font.size = Pt(11)
    run_int_title.font.bold = True
    run_int_title.font.color.rgb = RGBColor(71, 85, 105)

    nombres = [
        "1. DIAZ TORRES JAZMIN MONTSERRAT",
        "2. FIERRO MELÉNDEZ ERNESTO HATUEY",
        "3. HERNANDEZ GARCIA HECTOR GABRIEL",
        "4. MILLAN GUERRERO OSWALDO JOSUE"
    ]
    for nom in nombres:
        run_nom = integrantes_p.add_run(nom + "\n")
        run_nom.font.name = "Arial"
        run_nom.font.size = Pt(11)
        run_nom.font.bold = True
        run_nom.font.color.rgb = RGBColor(15, 23, 42)

    fecha_p = doc.add_paragraph()
    fecha_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fecha_p.paragraph_format.space_after = Pt(28)
    run_fecha = fecha_p.add_run("FECHA DE ENTREGA: 11 DE SEPTIEMBRE, 23:59")
    run_fecha.font.name = "Arial"
    run_fecha.font.size = Pt(10)
    run_fecha.font.italic = True
    run_fecha.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_page_break()

    # ---------------- SECCIÓN 1: VIDEO ----------------
    h1 = doc.add_heading("1. ENLACE PÚBLICO AL VIDEO DEMOSTRATIVO (YOUTUBE)", level=1)
    h1.paragraph_format.space_after = Pt(8)
    
    p_vid = doc.add_paragraph()
    p_vid.paragraph_format.line_spacing = 1.15
    run_url = p_vid.add_run("URL del Video (YouTube / Plataforma pública): ")
    run_url.font.bold = True
    p_vid.add_run("https://www.youtube.com/watch?v=TU_ENLACE_AQUI\n\n")
    
    p_vid_note = p_vid.add_run(
        "Nota: El video cumple con los lineamientos de la rúbrica oficial:\n"
        " • Duración máxima de 5 minutos y resolución mínima de 720p con audio explicativo claro.\n"
        " • Grabación de pantalla completa mostrando la fecha y hora visible del sistema en Windows.\n"
        " • Introducción con objetivo y presentación de los integrantes al inicio y fin de la ejecución.\n"
        " • Demostración de las reglas gramaticales para variables, constantes y expresiones.\n"
        " • Demostración de la captura de errores requeridos: identificador inválido, tipo no válido, falta de punto y coma, constante no inicializada y expresiones incompletas."
    )
    p_vid_note.font.size = Pt(9.5)
    p_vid_note.font.italic = True
    p_vid_note.font.color.rgb = RGBColor(71, 85, 105)

    # ---------------- SECCIÓN 2: INTRODUCCIÓN Y OBJETIVOS ----------------
    h2 = doc.add_heading("2. INTRODUCCIÓN Y OBJETIVOS", level=1)
    h2.paragraph_format.space_after = Pt(8)
    
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.line_spacing = 1.15
    p_intro.add_run(
        "El objetivo fundamental de esta práctica es diseñar e implementar la fase de Análisis Sintáctico (Parser) "
        "en la arquitectura de un compilador, centrado específicamente en la declaración y uso de variables y constantes, "
        "así como en el manejo robusto y pedagógico de errores sintácticos con recuperación en modo pánico.\n\n"
        "A diferencia del análisis léxico que opera secuencialmente sobre caracteres individuales para formar lexemas, "
        "el análisis sintáctico verifica que la estructura de la secuencia de tokens se apegue rigurosamente a una "
        "Gramática Libre de Contexto (GLC / CFG). Cuando una sentencia es válida, el analizador construye un Árbol de "
        "Sintaxis Abstracta (AST - Abstract Syntax Tree), el cual representa la estructura jerárquica del código fuente "
        "y servirá como entrada para las etapas posteriores de análisis semántico y generación de código intermedio."
    )

    # ---------------- SECCIÓN 3: DEFINICIÓN DE LA GRAMÁTICA ----------------
    h3 = doc.add_heading("3. DEFINICIÓN DE LA GRAMÁTICA FORMAL (EBNF / BNF)", level=1)
    h3.paragraph_format.space_after = Pt(8)

    p_ebnf_desc = doc.add_paragraph()
    p_ebnf_desc.add_run(
        "Se definió formalmente la gramática en Notación de Backus-Naur Extendida (EBNF) para soportar declaraciones de variables, "
        "constantes, asignaciones, expresiones aritméticas y lógicas con precedencia de operadores:"
    )

    ebnf_text = """programa             ::= ( declaracion | sentencia )* EOF

declaracion          ::= decl_variable | decl_constante

decl_variable        ::= tipo IDENTIFICADOR ( '=' expresion )? ( ',' IDENTIFICADOR ( '=' expresion )? )* ';'

decl_constante       ::= ( 'constante' | 'const' ) tipo IDENTIFICADOR '=' expresion ( ',' IDENTIFICADOR '=' expresion )* ';'

sentencia            ::= asignacion ';'
asignacion           ::= IDENTIFICADOR '=' expresion

tipo                 ::= 'entero' | 'flotante' | 'booleano' | 'cadena'
                       | 'int'    | 'float'    | 'bool'     | 'char' | 'string'

expresion            ::= expr_logica_or
expr_logica_or       ::= expr_logica_and ( ( '||' | 'or' ) expr_logica_and )*
expr_logica_and      ::= expr_igualdad ( ( '&&' | 'and' ) expr_igualdad )*
expr_igualdad        ::= expr_relacional ( ( '==' | '!=' ) expr_relacional )*
expr_relacional      ::= expr_aditiva ( ( '<' | '<=' | '>' | '>=' ) expr_aditiva )*
expr_aditiva         ::= expr_multiplicativa ( ( '+' | '-' ) expr_multiplicativa )*
expr_multiplicativa  ::= expr_unaria ( ( '*' | '/' | '%' ) expr_unaria )*
expr_unaria          ::= ( '+' | '-' | '!' | 'not' )? factor
factor               ::= IDENTIFICADOR
                       | LITERAL_ENTERO
                       | LITERAL_FLOTANTE
                       | LITERAL_CADENA
                       | LITERAL_BOOLEANO
                       | '(' expresion ')'
"""
    add_code_block(doc, ebnf_text, bg_hex="EFF6FF")

    # ---------------- SECCIÓN 4: CÓDIGO FUENTE ----------------
    h4 = doc.add_heading("4. CÓDIGO FUENTE DESARROLLADO", level=1)
    h4.paragraph_format.space_after = Pt(8)

    files_to_embed = [
        ("token_type.py", "Definición de tokens sintácticos y léxicos (Token y TokenType)."),
        ("lexer.py", "Analizador léxico con rastreo de línea, columna y detección de tokens inválidos."),
        ("ast_nodes.py", "Nodos del Árbol de Sintaxis Abstracta (AST) con visualización jerárquica."),
        ("error_handler.py", "Manejador centralizado de errores sintácticos con reglas específicas y sugerencias."),
        ("parser.py", "Analizador sintáctico descendente recursivo con recuperación en modo pánico."),
        ("test_cases.py", "Batería exhaustiva de casos de prueba válidos e inválidos."),
        ("main.py", "Ejecutable CLI con firmas obligatorias de integrantes al inicio y fin."),
        ("gui.py", "Interfaz gráfica interactiva en Tkinter para demostración visual de AST y errores.")
    ]

    current_dir = os.path.dirname(os.path.abspath(__file__))

    for filename, desc in files_to_embed:
        file_path = os.path.join(current_dir, filename)
        if os.path.exists(file_path):
            with open(file_path, "r", encoding="utf-8") as f:
                code_content = f.read()
            
            doc.add_heading(f"Archivo: {filename}", level=2)
            p_desc = doc.add_paragraph()
            p_desc.paragraph_format.space_after = Pt(4)
            r_d = p_desc.add_run(f"Descripción: {desc}")
            r_d.font.italic = True
            r_d.font.size = Pt(9.5)
            r_d.font.color.rgb = RGBColor(100, 116, 139)
            
            add_code_block(doc, code_content)

    # ---------------- SECCIÓN 5: RESULTADOS DE EJECUCIÓN ----------------
    doc.add_page_break()
    h5 = doc.add_heading("5. RESULTADOS DE EJECUCIÓN Y TABLAS DE DIAGNÓSTICO", level=1)
    h5.paragraph_format.space_after = Pt(8)

    p_res = doc.add_paragraph()
    p_res.add_run(
        "A continuación se presenta la salida íntegra de la ejecución de la suite de pruebas desde la terminal, "
        "demostrando la primera y última instrucción con los nombres de los integrantes en orden alfabético, "
        "los árboles sintácticos (AST) de casos válidos y las tablas de diagnóstico de errores sintácticos:"
    )

    add_code_block(doc, console_output, bg_hex="F1F5F9")

    # Guardar
    output_docx_path = os.path.join(current_dir, DOCX_NAME)
    doc.save(output_docx_path)
    print(f"✔ Documento Word guardado exitosamente en:\n  {output_docx_path}")

def main_generate():
    output = export_execution_log()
    generate_report_docx(output)

if __name__ == "__main__":
    main_generate()
